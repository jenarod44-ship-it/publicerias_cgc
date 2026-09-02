from datetime import datetime

from django.shortcuts import render
from django.utils import timezone
from datetime import datetime, timedelta


from .models import Servicio
from django.shortcuts import get_object_or_404, redirect, render

from .forms import ServicioDespachoForm
from .forms import ReposicionForm, ServicioDespachoForm
from .models import Servicio, TarifaChofer
from django.db import models
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from django.contrib.auth.decorators import login_required
from django.contrib.auth.decorators import user_passes_test

def puede_usar_despacho(user):
    return (
        user.is_superuser
        or user.groups.filter(
            name__in=["Administrador", "Coordinador", "Despachador"]
        ).exists()
    )

def obtener_datos_despacho(fecha):
    servicios = (
        Servicio.objects
        .filter(fecha=fecha)
        .select_related(
            "orden",
            "publicidad",
            "turno",
            "camion_contratado",
            "camion_operativo",
            "chofer",
            "telefono",
        )
        .order_by(
            "turno__hora_inicio",
            "camion_contratado__clave",
        )
    )

    total_servicios = servicios.count()
    programados = servicios.filter(
        estado="PROGRAMADO"
    ).count()

    en_servicio = servicios.filter(
        estado="EN_SERVICIO"
    ).count()

    realizados = servicios.filter(
        estado="REALIZADO"
    ).count()

    programados = servicios.filter(
        estado="PROGRAMADO"
    ).count()

    en_servicio = servicios.filter(
        estado="EN_SERVICIO"
    ).count()

    requieren_atencion = servicios.filter(
        estado__in=[
            "NO_SALIO",
            "PENDIENTE_REPOSICION",
        ]
    ).count()

    return {
        "fecha": fecha,
        "servicios": servicios,
        "total_servicios": total_servicios,
        "programados": programados,
        "en_servicio": en_servicio,
        "realizados": realizados,
        "requieren_atencion": requieren_atencion,
    }

@login_required
@user_passes_test(puede_usar_despacho)
def despacho_dia(request):
    fecha_texto = request.GET.get("fecha")

    if fecha_texto:
        try:
            fecha = datetime.strptime(
                fecha_texto,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            fecha = timezone.localdate()
    else:
        fecha = timezone.localdate()

    contexto = obtener_datos_despacho(fecha)

    return render(
        request,
        "core/despacho_dia.html",
        contexto,
    )

def puede_ver_dashboard(user):
    return (
        user.is_superuser
        or user.groups.filter(
            name__in=["Administrador", "Coordinador", "Gerencia"]
        ).exists()
    )

def puede_ver_reportes(user):
    return (
        user.is_superuser
        or user.groups.filter(
            name__in=["Administrador", "Coordinador", "Gerencia"]
        ).exists()
    )

@login_required
def reporte_despacho_dia(request):
    fecha_texto = request.GET.get("fecha")

    if fecha_texto:
        try:
            fecha = datetime.strptime(
                fecha_texto,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            fecha = timezone.localdate()
    else:
        fecha = timezone.localdate()

    contexto = obtener_datos_despacho(fecha)

    return render(
        request,
        "core/reporte_despacho_dia.html",
        contexto,
    )
@login_required
def editar_servicio(request, pk):
    servicio = get_object_or_404(Servicio, pk=pk)

    if request.method == "POST":
        form = ServicioDespachoForm(request.POST, instance=servicio)

        if form.is_valid():
            form.save()

            return redirect(
                f"/despacho/?fecha={servicio.fecha:%Y-%m-%d}"
            )
    else:
        form = ServicioDespachoForm(instance=servicio)

    contexto = {
        "servicio": servicio,
        "form": form,
    }

    return render(request, "core/editar_servicio.html", contexto)
@login_required
def programar_reposicion(request, pk):
    servicio_original = get_object_or_404(
        Servicio,
        pk=pk,
        estado="PENDIENTE_REPOSICION",
    )

    if servicio_original.reposiciones.exists():
        reposicion = servicio_original.reposiciones.first()

        return redirect(
            f"/despacho/?fecha={reposicion.fecha:%Y-%m-%d}"
        )

    if request.method == "POST":
        form = ReposicionForm(request.POST)

        if form.is_valid():
            reposicion = form.save(commit=False)

            reposicion.orden = servicio_original.orden
            reposicion.publicidad = servicio_original.publicidad
            reposicion.camion_contratado = servicio_original.camion_contratado
            reposicion.recorrido = servicio_original.recorrido
            reposicion.perifoneo = servicio_original.perifoneo

            reposicion.es_reposicion = True
            reposicion.servicio_original = servicio_original
            reposicion.estado = "PROGRAMADO"

            reposicion.save()

            return redirect(
                f"/despacho/?fecha={reposicion.fecha:%Y-%m-%d}"
            )
    else:
        form = ReposicionForm(
            initial={
                "turno": servicio_original.turno,
                "camion_operativo": servicio_original.camion_operativo
                or servicio_original.camion_contratado,
                "chofer": servicio_original.chofer,
            }
        )

    contexto = {
        "servicio_original": servicio_original,
        "form": form,
    }

    return render(
        request,
        "core/programar_reposicion.html",
        contexto,
    )

def obtener_datos_reporte_semanal(fecha_base):

    dias_desde_jueves = (fecha_base.weekday() - 3) % 7
    fecha_inicio = fecha_base - timedelta(days=dias_desde_jueves)
    fecha_fin = fecha_inicio + timedelta(days=6)

    tarifa_normal = (
        TarifaChofer.objects
        .filter(
            tipo="NORMAL",
            activa=True,
            fecha_inicio__lte=fecha_fin,
        )
        .filter(
            models.Q(fecha_fin__isnull=True) |
            models.Q(fecha_fin__gte=fecha_inicio)
        )
        .order_by("-fecha_inicio")
        .first()
    )

    tarifa_tercero = (
        TarifaChofer.objects
        .filter(
            tipo="TERCERO",
            activa=True,
            fecha_inicio__lte=fecha_fin,
        )
        .filter(
            models.Q(fecha_fin__isnull=True) |
            models.Q(fecha_fin__gte=fecha_inicio)
        )
        .order_by("-fecha_inicio")
        .first()
    )

    servicios = (
        Servicio.objects.filter(
            fecha__range=(fecha_inicio, fecha_fin),
            estado="REALIZADO",
            chofer__isnull=False,
        )
        .select_related(
            "chofer",
            "turno",
            "publicidad",
            "camion_operativo",
        )
        .order_by("chofer__nombre", "fecha", "turno")
    )

    resumen = {}

    for servicio in servicios:
        chofer = servicio.chofer

        if chofer.pk not in resumen:
            resumen[chofer.pk] = {
                "chofer": chofer,
                "primero": 0,
                "segundo": 0,
                "tercero": 0,
                "total": 0,
            }

        nombre_turno = str(servicio.turno).lower()

        if "primer" in nombre_turno:
            resumen[chofer.pk]["primero"] += 1
        elif "segundo" in nombre_turno:
            resumen[chofer.pk]["segundo"] += 1
        elif "tercer" in nombre_turno:
            resumen[chofer.pk]["tercero"] += 1

        resumen[chofer.pk]["total"] += 1

    importe_normal = tarifa_normal.importe if tarifa_normal else 0
    importe_tercero = tarifa_tercero.importe if tarifa_tercero else 0

    for fila in resumen.values():
        fila["importe_normal"] = importe_normal
        fila["importe_tercero"] = importe_tercero

        fila["total_pagar"] = (
            (fila["primero"] + fila["segundo"]) * importe_normal
            + fila["tercero"] * importe_tercero
        )

    total_general = sum(
        fila["total_pagar"]
        for fila in resumen.values()
    )

    return {
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "resumen": resumen.values(),
        "servicios": servicios,
        "tarifa_normal": tarifa_normal,
        "tarifa_tercero": tarifa_tercero,
        "total_general": total_general,
    }
@login_required
@user_passes_test(puede_ver_reportes)
def reporte_semanal_choferes(request):
    fecha_consulta = request.GET.get("fecha")

    if fecha_consulta:
        try:
            fecha_base = datetime.strptime(
                fecha_consulta,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            fecha_base = timezone.localdate()
    else:
        fecha_base = timezone.localdate()

    contexto = obtener_datos_reporte_semanal(fecha_base)

    contexto["fecha_base"] = fecha_base

    return render(
        request,
        "core/reporte_semanal_choferes.html",
        contexto,
    )
@login_required
@user_passes_test(puede_ver_reportes)
def exportar_reporte_semanal_excel(request):
    fecha_consulta = request.GET.get("fecha")

    if fecha_consulta:
        try:
            fecha_base = datetime.strptime(
                fecha_consulta,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            fecha_base = timezone.localdate()
    else:
        fecha_base = timezone.localdate()

    datos = obtener_datos_reporte_semanal(fecha_base)

    fecha_inicio = datos["fecha_inicio"]
    fecha_fin = datos["fecha_fin"]
    resumen = datos["resumen"]
    servicios = datos["servicios"]
    total_general = datos["total_general"]

    libro = Workbook()

    # =========================================================
    # HOJA 1 - RESUMEN SEMANAL
    # =========================================================
    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen semanal"

    hoja_resumen.merge_cells("A1:H1")
    hoja_resumen["A1"] = "Comunicadores Gráficos Creativos"
    hoja_resumen["A1"].font = Font(bold=True, size=14)
    hoja_resumen["A1"].alignment = Alignment(horizontal="center")

    hoja_resumen.merge_cells("A2:H2")
    hoja_resumen["A2"] = "Acumulado semanal de recorridos"
    hoja_resumen["A2"].font = Font(bold=True, size=12)
    hoja_resumen["A2"].alignment = Alignment(horizontal="center")

    hoja_resumen.merge_cells("A3:H3")
    hoja_resumen["A3"] = (
        f"Periodo: {fecha_inicio.strftime('%d/%m/%Y')} "
        f"al {fecha_fin.strftime('%d/%m/%Y')}"
    )
    hoja_resumen["A3"].alignment = Alignment(horizontal="center")

    encabezados = [
        "Chofer",
        "Primer turno",
        "Segundo turno",
        "Tercer turno",
        "Total recorridos",
        "Tarifa normal",
        "Tarifa tercero",
        "Total a pagar",
    ]

    fila_encabezados = 5

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja_resumen.cell(
            row=fila_encabezados,
            column=columna,
            value=encabezado,
        )
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")

    fila = fila_encabezados + 1

    for registro in resumen:
        hoja_resumen.cell(fila, 1, str(registro["chofer"]))
        hoja_resumen.cell(fila, 2, registro["primero"])
        hoja_resumen.cell(fila, 3, registro["segundo"])
        hoja_resumen.cell(fila, 4, registro["tercero"])
        hoja_resumen.cell(fila, 5, registro["total"])

        celda_normal = hoja_resumen.cell(
            fila,
            6,
            float(registro["importe_normal"]),
        )
        celda_normal.number_format = '$#,##0.00'

        celda_tercero = hoja_resumen.cell(
            fila,
            7,
            float(registro["importe_tercero"]),
        )
        celda_tercero.number_format = '$#,##0.00'

        celda_total = hoja_resumen.cell(
            fila,
            8,
            float(registro["total_pagar"]),
        )
        celda_total.number_format = '$#,##0.00'

        fila += 1

    hoja_resumen.cell(
        fila + 1,
        7,
        "TOTAL GENERAL:",
    ).font = Font(bold=True)

    celda_total_general = hoja_resumen.cell(
        fila + 1,
        8,
        float(total_general),
    )
    celda_total_general.font = Font(bold=True)
    celda_total_general.number_format = '$#,##0.00'

    anchos_resumen = {
        "A": 28,
        "B": 14,
        "C": 15,
        "D": 14,
        "E": 16,
        "F": 15,
        "G": 15,
        "H": 16,
    }

    for columna, ancho in anchos_resumen.items():
        hoja_resumen.column_dimensions[columna].width = ancho

    hoja_resumen.freeze_panes = "A6"

    # =========================================================
    # HOJA 2 - DETALLE DE RECORRIDOS
    # =========================================================
    hoja_detalle = libro.create_sheet("Detalle recorridos")

    hoja_detalle.merge_cells("A1:G1")
    hoja_detalle["A1"] = "Detalle de recorridos realizados"
    hoja_detalle["A1"].font = Font(bold=True, size=14)
    hoja_detalle["A1"].alignment = Alignment(horizontal="center")

    hoja_detalle.merge_cells("A2:G2")
    hoja_detalle["A2"] = (
        f"Periodo: {fecha_inicio.strftime('%d/%m/%Y')} "
        f"al {fecha_fin.strftime('%d/%m/%Y')}"
    )
    hoja_detalle["A2"].alignment = Alignment(horizontal="center")

    encabezados_detalle = [
        "Fecha",
        "Chofer",
        "Publicidad",
        "Camión",
        "Turno",
        "Salida",
        "Regreso",
    ]

    for columna, encabezado in enumerate(
        encabezados_detalle,
        start=1,
    ):
        celda = hoja_detalle.cell(
            row=4,
            column=columna,
            value=encabezado,
        )
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")

    fila = 5

    for servicio in servicios:
        celda_fecha = hoja_detalle.cell(
            fila,
            1,
            servicio.fecha,
        )
        celda_fecha.number_format = "dd/mm/yyyy"

        hoja_detalle.cell(
            fila,
            2,
            str(servicio.chofer),
        )

        hoja_detalle.cell(
            fila,
            3,
            str(servicio.publicidad),
        )

        camion = (
            servicio.camion_operativo
            or servicio.camion_contratado
        )

        hoja_detalle.cell(
            fila,
            4,
            str(camion) if camion else "",
        )

        hoja_detalle.cell(
            fila,
            5,
            str(servicio.turno),
        )

        celda_salida = hoja_detalle.cell(
            fila,
            6,
            servicio.hora_salida,
        )
        if servicio.hora_salida:
            celda_salida.number_format = "hh:mm"

        celda_regreso = hoja_detalle.cell(
            fila,
            7,
            servicio.hora_regreso,
        )
        if servicio.hora_regreso:
            celda_regreso.number_format = "hh:mm"

        fila += 1

    anchos_detalle = {
        "A": 13,
        "B": 28,
        "C": 30,
        "D": 12,
        "E": 20,
        "F": 12,
        "G": 12,
    }

    for columna, ancho in anchos_detalle.items():
        hoja_detalle.column_dimensions[columna].width = ancho

    hoja_detalle.freeze_panes = "A5"

    nombre_archivo = (
        f"reporte_semanal_"
        f"{fecha_inicio.strftime('%Y%m%d')}_"
        f"{fecha_fin.strftime('%Y%m%d')}.xlsx"
    )

    respuesta = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    respuesta["Content-Disposition"] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    libro.save(respuesta)

    return respuesta

def obtener_datos_reporte_servicios(fecha_desde, fecha_hasta):
    servicios = (
        Servicio.objects
        .filter(
            fecha__range=(fecha_desde, fecha_hasta)
        )
        .select_related(
            "orden",
            "publicidad",
            "turno",
            "camion_contratado",
            "camion_operativo",
            "chofer",
            "telefono",
            "motivo_no_salida",
        )
        .order_by(
            "fecha",
            "turno__hora_inicio",
            "camion_contratado__clave",
        )
    )

    return {
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "servicios": servicios,
        "total_servicios": servicios.count(),
    }
@login_required
@user_passes_test(puede_ver_reportes)
def reporte_servicios(request):
    fecha_desde_texto = request.GET.get("desde")
    fecha_hasta_texto = request.GET.get("hasta")

    hoy = timezone.localdate()

    try:
        fecha_desde = (
            datetime.strptime(fecha_desde_texto, "%Y-%m-%d").date()
            if fecha_desde_texto
            else hoy
        )
    except ValueError:
        fecha_desde = hoy

    try:
        fecha_hasta = (
            datetime.strptime(fecha_hasta_texto, "%Y-%m-%d").date()
            if fecha_hasta_texto
            else fecha_desde
        )
    except ValueError:
        fecha_hasta = fecha_desde

    if fecha_hasta < fecha_desde:
        fecha_hasta = fecha_desde

    contexto = obtener_datos_reporte_servicios(
        fecha_desde,
        fecha_hasta,
    )

    return render(
        request,
        "core/reporte_servicios.html",
        contexto,
    )

@login_required
@user_passes_test(puede_ver_reportes)
def exportar_reporte_servicios_excel(request):
    fecha_desde_texto = request.GET.get("desde")
    fecha_hasta_texto = request.GET.get("hasta")

    hoy = timezone.localdate()

    try:
        fecha_desde = (
            datetime.strptime(fecha_desde_texto, "%Y-%m-%d").date()
            if fecha_desde_texto
            else hoy
        )
    except ValueError:
        fecha_desde = hoy

    try:
        fecha_hasta = (
            datetime.strptime(fecha_hasta_texto, "%Y-%m-%d").date()
            if fecha_hasta_texto
            else fecha_desde
        )
    except ValueError:
        fecha_hasta = fecha_desde

    if fecha_hasta < fecha_desde:
        fecha_hasta = fecha_desde

    datos = obtener_datos_reporte_servicios(
        fecha_desde,
        fecha_hasta,
    )

    servicios = datos["servicios"]

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Servicios"

    hoja.merge_cells("A1:I1")
    hoja["A1"] = "Comunicadores Gráficos Creativos"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A1"].alignment = Alignment(horizontal="center")

    hoja.merge_cells("A2:I2")
    hoja["A2"] = "Reporte de Servicios"
    hoja["A2"].font = Font(bold=True, size=12)
    hoja["A2"].alignment = Alignment(horizontal="center")

    hoja.merge_cells("A3:I3")
    hoja["A3"] = (
        f"Periodo: {fecha_desde.strftime('%d/%m/%Y')} "
        f"al {fecha_hasta.strftime('%d/%m/%Y')}"
    )
    hoja["A3"].alignment = Alignment(horizontal="center")

    encabezados = [
        "Fecha",
        "Publicidad",
        "Camión",
        "Chofer",
        "Turno",
        "Recorrido",
        "Salida",
        "Regreso",
        "Estado",
    ]

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(
            row=5,
            column=columna,
            value=encabezado,
        )
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")

    fila = 6

    for servicio in servicios:
        celda_fecha = hoja.cell(
            fila,
            1,
            servicio.fecha,
        )
        celda_fecha.number_format = "dd/mm/yyyy"

        hoja.cell(
            fila,
            2,
            str(servicio.publicidad),
        )

        camion = (
            servicio.camion_operativo
            or servicio.camion_contratado
        )

        hoja.cell(
            fila,
            3,
            str(camion) if camion else "",
        )

        hoja.cell(
            fila,
            4,
            str(servicio.chofer) if servicio.chofer else "",
        )

        hoja.cell(
            fila,
            5,
            str(servicio.turno),
        )

        hoja.cell(
            fila,
            6,
            servicio.recorrido,
        )

        celda_salida = hoja.cell(
            fila,
            7,
            servicio.hora_salida,
        )
        if servicio.hora_salida:
            celda_salida.number_format = "hh:mm"

        celda_regreso = hoja.cell(
            fila,
            8,
            servicio.hora_regreso,
        )
        if servicio.hora_regreso:
            celda_regreso.number_format = "hh:mm"

        hoja.cell(
            fila,
            9,
            servicio.get_estado_display(),
        )

        fila += 1

    hoja.cell(
        fila + 1,
        8,
        "TOTAL SERVICIOS:",
    ).font = Font(bold=True)

    hoja.cell(
        fila + 1,
        9,
        datos["total_servicios"],
    ).font = Font(bold=True)

    anchos = {
        "A": 13,
        "B": 30,
        "C": 12,
        "D": 28,
        "E": 20,
        "F": 45,
        "G": 12,
        "H": 12,
        "I": 24,
    }

    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    hoja.freeze_panes = "A6"
    hoja.auto_filter.ref = f"A5:I{fila - 1}"

    nombre_archivo = (
        f"reporte_servicios_"
        f"{fecha_desde.strftime('%Y%m%d')}_"
        f"{fecha_hasta.strftime('%Y%m%d')}.xlsx"
    )

    respuesta = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    respuesta["Content-Disposition"] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    libro.save(respuesta)

    return respuesta


@login_required
@user_passes_test(puede_ver_dashboard)
def dashboard_gerencial(request):
    hoy = timezone.localdate()

    dias_desde_jueves = (hoy.weekday() - 3) % 7
    fecha_inicio = hoy - timedelta(days=dias_desde_jueves)
    fecha_fin = fecha_inicio + timedelta(days=6)

    servicios = Servicio.objects.filter(
        fecha__range=(fecha_inicio, fecha_fin)
    )

    total_servicios = servicios.count()

    realizados = servicios.filter(
        estado="REALIZADO"
    ).count()

    programados = servicios.filter(
        estado="PROGRAMADO"
    ).count()

    en_servicio = servicios.filter(
        estado="EN_SERVICIO"
    ).count()

    requieren_atencion = servicios.filter(
        estado__in=[
            "NO_SALIO",
            "PENDIENTE_REPOSICION",
        ]
    ).count()

    servicios_atencion = (
        servicios
        .filter(
            estado__in=[
                "NO_SALIO",
                "PENDIENTE_REPOSICION",
            ]
        )
        .select_related(
            "publicidad",
            "turno",
            "camion_operativo",
            "camion_contratado",
            "chofer",
            "motivo_no_salida",
        )
        .order_by(
            "fecha",
            "turno__hora_inicio",
        )
    )

    datos_semanales = obtener_datos_reporte_semanal(hoy)

    contexto = {
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "total_servicios": total_servicios,
        "realizados": realizados,
        "programados": programados,
        "en_servicio": en_servicio,
        "requieren_atencion": requieren_atencion,
        "total_pagar": datos_semanales["total_general"],
        "servicios_atencion": servicios_atencion,
        "resumen_choferes": datos_semanales["resumen"],
    }

    return render(
        request,
        "core/dashboard_gerencial.html",
        contexto,
    )